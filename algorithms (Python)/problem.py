"""
Optimization Benchmark Problem Harness (Example Benchmark)
==========================================================
Provides the problem definition and standardized evaluation harness:
- Problem search space bounds (lb, ub) and dimension (dim)
- Objective function (fitness evaluator)
- Experiment runner with multi-run statistical evaluation, figure generation,
  and comprehensive multi-sheet Excel export.

This file serves as a plug-and-play benchmark example. Users can easily replace
the objective function and search bounds with their own target problem.
"""

import os
import time
import warnings
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

# Safely suppress convergence warnings without requiring scikit-learn
warnings.filterwarnings("ignore")
try:
    from sklearn.exceptions import ConvergenceWarning
    warnings.filterwarnings("ignore", category=ConvergenceWarning)
except ImportError:
    pass

# =========================================================
# Problem Identification Metadata
# =========================================================
PROBLEM_NAME = "Example Benchmark Problem"
PROBLEM_GOAL = "Minimization"
KNOWN_OPTIMUM = None

# =========================================================
# Experimental Data — RTC France Si Solar Cell
# =========================================================
V_exp = np.array([
    -0.2057, -0.1291, -0.0588,  0.0057,  0.0646,  0.1185,
     0.1678,  0.2132,  0.2545,  0.2924,  0.3269,  0.3585,
     0.3873,  0.4137,  0.4373,  0.4590,  0.4784,  0.4960,
     0.5119,  0.5265,  0.5398,  0.5521,  0.5633,  0.5736,
     0.5833,  0.5900
])
I_exp = np.array([
     0.7640,  0.7620,  0.7605,  0.7605,  0.7600,  0.7590,
     0.7570,  0.7570,  0.7555,  0.7540,  0.7505,  0.7465,
     0.7385,  0.7280,  0.7065,  0.6755,  0.6320,  0.5730,
     0.4990,  0.4130,  0.3165,  0.2120,  0.1035, -0.0100,
    -0.1230, -0.2100
])

# =========================================================
# Physical Constants
# =========================================================
k  = 1.3806503e-23
q  = 1.60217646e-19
T  = 273.15 + 33
Vt = k * T / q

# =========================================================
# Bounds — MTDM
# Params: [Iph, Isd1, Isd2, Isd3, Rs, Rsh, n1, n2, n3, Rsm]
# =========================================================
lb = np.array([0.0,  0.0,   0.0,   0.0,   0.0,   0.0,  1.0,  1.0,  1.0,  0.0])
ub = np.array([1.0,  1e-6,  1e-6,  1e-6,  0.5,  100.0, 2.0,  2.0,  2.0,  0.5])
dim = 10

# =========================================================
# Modified Triple Diode Model — Vectorized Coupled Newton-Raphson
# =========================================================
def modified_triple_diode_model(V, params, max_iter=150, tol=1e-12):
    """
    MTDM equation (implicit in both I and I_D3):

      I = Iph
          - Isd1*(exp((V + I*Rs)/(n1*Vt)) - 1)
          - Isd2*(exp((V + I*Rs)/(n2*Vt)) - 1)
          - Isd3*(exp((V + I*Rs - Rsm*I_D3)/(n3*Vt)) - 1)
          - (V + I*Rs)/Rsh

      I_D3 = Isd3*(exp((V + I*Rs - Rsm*I_D3)/(n3*Vt)) - 1)

    Solved via a vectorized 2x2 Newton-Raphson on the unknowns [I, I_D3]
    simultaneously (closed-form 2x2 linear solve at every iteration).
    """
    Iph, Isd1, Isd2, Isd3, Rs, Rsh, n1, n2, n3, Rsm = params

    Rsh = Rsh if abs(Rsh) > 1e-10 else 1e-10

    n1Vt = n1 * Vt
    n2Vt = n2 * Vt
    n3Vt = n3 * Vt

    I   = np.full_like(V, Iph, dtype=np.float64)
    ID3 = np.zeros_like(V, dtype=np.float64)

    for _ in range(max_iter):
        VIRs = V + I * Rs

        arg1 = np.clip(VIRs / n1Vt, -500.0, 500.0)
        arg2 = np.clip(VIRs / n2Vt, -500.0, 500.0)
        arg3 = np.clip((VIRs - Rsm * ID3) / n3Vt, -500.0, 500.0)

        ex1 = np.exp(arg1)
        ex2 = np.exp(arg2)
        ex3 = np.exp(arg3)

        # Residuals
        F1 = (
            I - Iph
            + Isd1 * (ex1 - 1.0)
            + Isd2 * (ex2 - 1.0)
            + Isd3 * (ex3 - 1.0)
            + VIRs / Rsh
        )
        F2 = ID3 - Isd3 * (ex3 - 1.0)

        # Jacobian entries
        dF1_dI   = (
            1.0
            + (Isd1 * Rs / n1Vt) * ex1
            + (Isd2 * Rs / n2Vt) * ex2
            + (Isd3 * Rs / n3Vt) * ex3
            + Rs / Rsh
        )
        dF1_dID3 = -(Isd3 * Rsm / n3Vt) * ex3
        dF2_dI   = -(Isd3 * Rs / n3Vt) * ex3
        dF2_dID3 = 1.0 + (Isd3 * Rsm / n3Vt) * ex3

        # Closed-form 2x2 linear solve: J * [dI; dID3] = [F1; F2]
        det = dF1_dI * dF2_dID3 - dF1_dID3 * dF2_dI
        det = np.where(np.abs(det) < 1e-30, 1e-30, det)

        dI   = (F1 * dF2_dID3 - F2 * dF1_dID3) / det
        dID3 = (dF1_dI * F2 - dF2_dI * F1) / det

        I   -= dI
        ID3 -= dID3

        if np.max(np.abs(dI)) < tol and np.max(np.abs(dID3)) < tol:
            break

    return I


# Backward-compatible alias, in case other code imports the old name
triple_diode_model = modified_triple_diode_model

# =========================================================
# Objective Function — RMSE
# =========================================================
def objective_function(params):
    try:
        params = np.clip(params, lb, ub)
        I_calc = modified_triple_diode_model(V_exp, params)
        if not np.all(np.isfinite(I_calc)):
            return 1e9
        return float(np.sqrt(np.mean((I_exp - I_calc) ** 2)))
    except Exception:
        return 1e9

# =========================================================
# Excel Export
# =========================================================
def save_comprehensive_excel(all_convergence, all_rmse, all_best_params,
                              V_exp, I_exp, best_parameters,
                              max_iter, n_runs, all_times,
                              algo_name, full_name=None,
                              filename='Results.xlsx'):
    wb = Workbook()
    wb.remove(wb.active)

    best_idx  = int(np.argmin(all_rmse))
    worst_idx = int(np.argmax(all_rmse))
    ranks     = np.argsort(np.argsort(all_rmse)) + 1

    C_HEADER  = "1F4E79"
    C_GREEN_H = "375623"
    C_RED_H   = "7B0000"
    C_BEST    = "E2EFDA"
    C_WORST   = "FFE0E0"
    C_ALT     = "F2F7FC"
    C_WHITE   = "FFFFFF"

    thin   = Side(style='thin', color="AAAAAA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def hdr(ws, n, color=C_HEADER):
        for c in range(1, n + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill      = PatternFill("solid", fgColor=color)
            cell.font      = Font(bold=True, color="FFFFFF", size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border    = border
        ws.row_dimensions[1].height = 18

    def style_data(ws, n_rows, n_cols, start_row=2, best_row=None, worst_row=None):
        for r in range(start_row, start_row + n_rows):
            fc = C_WHITE
            if best_row and r == best_row:
                fc = C_BEST
            elif worst_row and r == worst_row:
                fc = C_WORST
            elif (r - start_row) % 2 == 1:
                fc = C_ALT
            for c in range(1, n_cols + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill      = PatternFill("solid", fgColor=fc)
                cell.alignment = Alignment(horizontal="center")
                cell.border    = border

    def autofit(ws):
        for col in ws.columns:
            ml = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(ml + 4, 30)

    # Sheet 1 — Summary
    ws1 = wb.create_sheet("Summary")
    h1  = ['Run', 'Rank', 'RMSE', 'Time (s)', 'Iph (A)', 'Isd1 (A)', 'Isd2 (A)', 'Isd3 (A)',
           'Rs (Ω)', 'Rsh (Ω)', 'n1', 'n2', 'n3', 'Rsm (Ω)', 'Note']
    ws1.append(h1)
    for i in range(n_runs):
        note = '★ BEST' if i == best_idx else ('✗ WORST' if i == worst_idx else '')
        p = all_best_params[i]
        ws1.append([
            i + 1, int(ranks[i]), round(float(all_rmse[i]), 10), round(float(all_times[i]), 2),
            round(float(p[0]), 8), f"{float(p[1]):.4e}", f"{float(p[2]):.4e}", f"{float(p[3]):.4e}",
            round(float(p[4]), 8), round(float(p[5]), 6),
            round(float(p[6]), 8), round(float(p[7]), 8), round(float(p[8]), 8),
            round(float(p[9]), 8), note
        ])
    hdr(ws1, len(h1))
    style_data(ws1, n_runs, len(h1), best_row=best_idx + 2, worst_row=worst_idx + 2)
    autofit(ws1)
    ws1.freeze_panes = 'A2'

    # Sheet 2 — Convergence
    ws2 = wb.create_sheet("Convergence_All")
    ws2.append(['Iteration'] + [f'Run_{i+1:02d}' for i in range(n_runs)])
    for it in range(max_iter):
        ws2.append([it + 1] + [round(float(all_convergence[r][it]), 10) for r in range(n_runs)])
    hdr(ws2, n_runs + 1)
    for r in range(1, max_iter + 2):
        ws2.cell(row=r, column=best_idx + 2).fill = PatternFill("solid", fgColor="D9EAD3")
        ws2.cell(row=r, column=worst_idx + 2).fill = PatternFill("solid", fgColor="FFE0E0")
    ws2.cell(row=1, column=best_idx + 2).fill = PatternFill("solid", fgColor=C_GREEN_H)
    ws2.cell(row=1, column=worst_idx + 2).fill = PatternFill("solid", fgColor=C_RED_H)
    autofit(ws2)
    ws2.freeze_panes = 'B2'

    # Sheet 3 — IV / PV
    ws3 = wb.create_sheet("Best_Run_IV_PV")
    V_sm = np.linspace(float(np.min(V_exp)) - 0.05, float(np.max(V_exp)) + 0.05, 300)
    I_sm = modified_triple_diode_model(V_sm, best_parameters)
    P_sm = V_sm * I_sm
    I_c  = modified_triple_diode_model(V_exp, best_parameters)
    P_e  = V_exp * I_exp
    P_c  = V_exp * I_c
    h3 = ['V_smooth (V)', 'I_model (A)', 'P_model (W)', '', 'V_exp (V)', 'I_exp (A)',
          'P_exp (W)', 'I_model_at_exp (A)', 'P_model_at_exp (W)']
    ws3.append(h3)
    for ri in range(max(len(V_sm), len(V_exp))):
        row = []
        if ri < len(V_sm):
            row += [round(float(V_sm[ri]), 6), round(float(I_sm[ri]), 8), round(float(P_sm[ri]), 8)]
        else:
            row += ['', '', '']
        row.append('')
        if ri < len(V_exp):
            row += [round(float(V_exp[ri]), 6), round(float(I_exp[ri]), 8), round(float(P_e[ri]), 8),
                    round(float(I_c[ri]), 8), round(float(P_c[ri]), 8)]
        ws3.append(row)
    hdr(ws3, len(h3))
    style_data(ws3, max(len(V_sm), len(V_exp)), len(h3))
    autofit(ws3)
    ws3.freeze_panes = 'A2'

    # Sheet 4 — Exp vs Pred
    ws4 = wb.create_sheet("Exp_vs_Pred")
    h4  = ['#', 'V (V)', 'I_exp (A)', 'I_pred (A)', 'Abs Error (A)', 'Error %',
           'P_exp (W)', 'P_pred (W)', 'P Error (W)']
    ws4.append(h4)
    I_c2 = modified_triple_diode_model(V_exp, best_parameters)
    P_e2 = V_exp * I_exp
    P_c2 = V_exp * I_c2
    for i in range(len(V_exp)):
        err   = float(I_exp[i] - I_c2[i])
        ep    = abs(err) / abs(float(I_exp[i])) * 100 if I_exp[i] != 0 else 0.0
        perr  = float(P_e2[i] - P_c2[i])
        ws4.append([i + 1, round(float(V_exp[i]), 6), round(float(I_exp[i]), 8),
                    round(float(I_c2[i]), 8), round(abs(err), 10), round(ep, 6),
                    round(float(P_e2[i]), 8), round(float(P_c2[i]), 8), round(abs(perr), 10)])
    mae_i2 = float(np.mean(np.abs(I_exp - I_c2)))
    mae_p2 = float(np.mean(np.abs(P_e2 - P_c2)))
    ws4.append([])
    ws4.append(['MAE', '', '', '', round(mae_i2, 10), '', '', '', round(mae_p2, 10)])
    ws4.append(['RMSE', '', '', '', round(float(all_rmse[best_idx]), 10), '', '', '', ''])
    ws4.append(['Max Err', '', '', '', round(float(np.max(np.abs(I_exp - I_c2))), 10), '', '', '', ''])
    hdr(ws4, len(h4))
    style_data(ws4, len(V_exp), len(h4))
    autofit(ws4)
    ws4.freeze_panes = 'A2'

    # Sheet 5 — Statistics
    ws5 = wb.create_sheet("Statistics")
    ws5.append(['Metric', 'Value'])
    hdr(ws5, 2, color="4A235A")
    display_title = full_name if full_name else algo_name
    stats = [
        ('Algorithm', display_title),
        ('Problem', PROBLEM_NAME),
        ('Runs', n_runs),
        ('Iterations/Run', max_iter),
        ('Agents', len(all_best_params[0]) if len(all_best_params) > 0 else 30),
        ('', ''),
        ('── RMSE Statistics ──', ''),
        ('Min (Best)',   round(float(np.min(all_rmse)), 10)),
        ('Max (Worst)',  round(float(np.max(all_rmse)), 10)),
        ('Mean',         round(float(np.mean(all_rmse)), 10)),
        ('Median',       round(float(np.median(all_rmse)), 10)),
        ('Std Dev',      round(float(np.std(all_rmse)), 10)),
        ('Variance',     round(float(np.var(all_rmse)), 10)),
        ('', ''),
        ('── Timing Statistics (s) ──', ''),
        ('Min',          round(float(np.min(all_times)), 4)),
        ('Max',          round(float(np.max(all_times)), 4)),
        ('Mean',         round(float(np.mean(all_times)), 4)),
        ('Median',       round(float(np.median(all_times)), 4)),
        ('Std Dev',      round(float(np.std(all_times)), 4)),
        ('Variance',     round(float(np.var(all_times)), 6)),
        ('Total',        round(float(np.sum(all_times)), 2)),
        ('', ''),
        ('── Best Run Parameters (MTDM) ──', ''),
        ('Run #', best_idx + 1),
        ('Iph (A)',   round(float(best_parameters[0]), 8)),
        ('Isd1 (A)', f"{float(best_parameters[1]):.6e}"),
        ('Isd2 (A)', f"{float(best_parameters[2]):.6e}"),
        ('Isd3 (A)', f"{float(best_parameters[3]):.6e}"),
        ('Rs (Ω)',   round(float(best_parameters[4]), 8)),
        ('Rsh (Ω)',  round(float(best_parameters[5]), 6)),
        ('n1',       round(float(best_parameters[6]), 8)),
        ('n2',       round(float(best_parameters[7]), 8)),
        ('n3',       round(float(best_parameters[8]), 8)),
        ('Rsm (Ω)',  round(float(best_parameters[9]), 8)),
        ('', ''),
        ('── Error Metrics ──', ''),
        ('RMSE (A)',          round(float(all_rmse[best_idx]), 10)),
        ('MAE (A)',           round(mae_i2, 10)),
        ('Max Abs Error (A)', round(float(np.max(np.abs(I_exp - I_c2))), 10)),
    ]
    for name, val in stats:
        ri = ws5.max_row + 1
        ws5.append([name, val])
        if '──' in str(name):
            for c in range(1, 3):
                cell = ws5.cell(row=ri, column=c)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2E4057")
                cell.alignment = Alignment(horizontal="left")
        elif name:
            ws5.cell(row=ri, column=1).font = Font(bold=True)
            ws5.cell(row=ri, column=2).alignment = Alignment(horizontal="center")
    autofit(ws5)
    ws5.freeze_panes = 'A2'

    wb.save(filename)
    print(f"\n{'='*60}")
    print(f"  Excel saved: {filename}")
    print(f"  Sheets: Summary | Convergence_All | Best_Run_IV_PV | Exp_vs_Pred | Statistics")
    print(f"  Best  Run: #{best_idx+1:02d}  RMSE = {all_rmse[best_idx]:.8e}")
    print(f"  Worst Run: #{worst_idx+1:02d}  RMSE = {all_rmse[worst_idx]:.8e}")
    print(f"{'='*60}")


# =========================================================
# Main Experiment Runner
# =========================================================
def run_tdm_experiment(algo_name, algo_func, full_name=None,
                       n_agents=30, max_iter=500, n_runs=4,
                       figures_dir="results/figures",
                       excel_dir="results/excel"):
    """
    Executes an n_runs independent experiment on the Modified Triple Diode
    Model (MTDM) parameter estimation benchmark using algo_func, generates
    all 4 figures, computes full statistics, and exports the 5-sheet Excel
    report.
    """
    display_title = full_name if full_name else algo_name
    os.makedirs(figures_dir, exist_ok=True)
    os.makedirs(excel_dir, exist_ok=True)

    print("=" * 70)
    print(f"  {display_title} — Benchmark Optimization Experiment")
    print(f"  Problem: {PROBLEM_NAME} (dim={dim})")
    print("=" * 70)
    print(f"\nRunning {n_runs} independent runs (Agents: {n_agents}, Iterations: {max_iter})...\n")

    all_best_params = []
    all_rmse        = []
    all_convergence = []
    all_times       = []

    for run in range(n_runs):
        t0 = time.perf_counter()
        print(f"  Run {run+1:>2}/{n_runs} ...", end=" ", flush=True)

        # Call optimizer function
        raw_res = algo_func(n_agents, max_iter, lb, ub, dim, objective_function)

        elapsed = time.perf_counter() - t0

        # Robust unpacking: handles (score, pos, curve) vs (pos, score, curve)
        if len(raw_res) == 3:
            r0, r1, r2 = raw_res
            if not np.isscalar(r0) and np.isscalar(r1):
                best_params = np.asarray(r0, dtype=float)
                best_rmse   = float(r1)
                convergence = np.asarray(r2, dtype=float)
            else:
                best_rmse   = float(r0)
                best_params = np.asarray(r1, dtype=float)
                convergence = np.asarray(r2, dtype=float)
        else:
            raise ValueError(f"Expected 3 return values from {algo_name}, got {len(raw_res)}")

        all_best_params.append(best_params)
        all_rmse.append(best_rmse)
        all_convergence.append(convergence)
        all_times.append(elapsed)

        print(f"RMSE = {best_rmse:.12e}  |  time = {elapsed:.1f}s")

    # =========================================================
    # Statistical Analysis
    # =========================================================
    all_rmse        = np.array(all_rmse)
    all_best_params = np.array(all_best_params)
    all_times       = np.array(all_times)

    min_rmse      = np.min(all_rmse)
    max_rmse      = np.max(all_rmse)
    mean_rmse     = np.mean(all_rmse)
    median_rmse   = np.median(all_rmse)
    std_rmse      = np.std(all_rmse)
    variance_rmse = np.var(all_rmse)

    best_idx         = int(np.argmin(all_rmse))
    best_parameters  = all_best_params[best_idx]
    best_convergence = all_convergence[best_idx]

    print("\n" + "=" * 70)
    print(f"  STATISTICAL RESULTS ({algo_name}):")
    print("-" * 70)
    print(f"  Min RMSE:      {min_rmse:.6e}")
    print(f"  Mean RMSE:     {mean_rmse:.6e}")
    print(f"  Max RMSE:      {max_rmse:.6e}")
    print(f"  Median RMSE:   {median_rmse:.6e}")
    print(f"  Std RMSE:      {std_rmse:.6e}")
    print(f"  Variance RMSE: {variance_rmse:.1e}")

    print("\n" + "=" * 70)
    print(f"  OPTIMAL PARAMETERS — Best Run (Run {best_idx+1}):")
    print("-" * 70)
    print(f"  Iph   (Photocurrent):       {best_parameters[0]:.8f} A")
    print(f"  Isd1  (Sat. cur. diode 1):  {best_parameters[1]:.6e} A")
    print(f"  Isd2  (Sat. cur. diode 2):  {best_parameters[2]:.6e} A")
    print(f"  Isd3  (Sat. cur. diode 3):  {best_parameters[3]:.6e} A")
    print(f"  Rs    (Series res.):         {best_parameters[4]:.8f} Ω")
    print(f"  Rsh   (Shunt res.):          {best_parameters[5]:.6f} Ω")
    print(f"  n1    (Ideality factor 1):   {best_parameters[6]:.8f}")
    print(f"  n2    (Ideality factor 2):   {best_parameters[7]:.8f}")
    print(f"  n3    (Ideality factor 3):   {best_parameters[8]:.8f}")
    print(f"  Rsm   (Diode-3 branch res.): {best_parameters[9]:.8f} Ω")
    print("=" * 70)

    print("\n" + "=" * 70)
    print("  TIMING RESULTS:")
    print("-" * 70)
    print(f"  Total Time:     {np.sum(all_times):.2f}s")
    print(f"  Mean/Run:       {np.mean(all_times):.2f}s")
    print(f"  Min (Run {np.argmin(all_times)+1:02d}):  {np.min(all_times):.2f}s")
    print(f"  Max (Run {np.argmax(all_times)+1:02d}):  {np.max(all_times):.2f}s")
    print("=" * 70)

    # =========================================================
    # Experimental vs Predicted Table
    # =========================================================
    I_predicted_best = modified_triple_diode_model(V_exp, best_parameters)
    P_exp_arr        = V_exp * I_exp
    P_pred_arr       = V_exp * I_predicted_best
    error_current    = I_exp - I_predicted_best
    error_power      = P_exp_arr - P_pred_arr

    print("\n" + "=" * 70)
    print("  EXPERIMENTAL vs PREDICTED CURRENT")
    print("=" * 70)
    print(f"  {'#':<5} {'V (V)':<12} {'I_exp (A)':<13} {'I_pred (A)':<13} {'Error (A)':<14} {'Error %'}")
    print("  " + "-" * 65)
    for i in range(len(V_exp)):
        err    = I_exp[i] - I_predicted_best[i]
        err_pc = abs(err) / abs(I_exp[i]) * 100 if I_exp[i] != 0 else 0.0
        print(f"  {i:<5} {V_exp[i]:<12.4f} {I_exp[i]:<13.6f} "
              f"{I_predicted_best[i]:<13.6f} {err:<14.6e} {err_pc:.4f}")

    mae     = np.mean(np.abs(error_current))
    max_err = np.max(np.abs(error_current))
    print("  " + "-" * 65)
    print(f"  RMSE:               {all_rmse[best_idx]:.6e}")
    print(f"  MAE:                {mae:.6e}")
    print(f"  Max Absolute Error: {max_err:.6e}")
    print("=" * 70)

    # =========================================================
    # Smooth curves for plotting
    # =========================================================
    V_smooth = np.linspace(-0.3, 0.62, 300)
    I_smooth = modified_triple_diode_model(V_smooth, best_parameters)
    P_smooth = V_smooth * I_smooth

    # =========================================================
    # Figure 1 — Main Results (2×2)
    # =========================================================
    fig1, axes = plt.subplots(2, 2, figsize=(15, 10))

    axes[0, 0].semilogy(best_convergence, 'b-', linewidth=2)
    axes[0, 0].set_xlabel('Iteration', fontsize=12, fontweight='bold')
    axes[0, 0].set_ylabel('RMSE', fontsize=12, fontweight='bold')
    axes[0, 0].set_title(f'{algo_name} Convergence Curve (Best Run)', fontsize=14, fontweight='bold')
    axes[0, 0].grid(True, alpha=0.3)

    axes[0, 1].plot(V_smooth, I_smooth, 'r-', linewidth=2, label=f'{algo_name} - MTDM')
    axes[0, 1].plot(V_exp, I_exp, 'bo', markersize=6, label='Experimental')
    axes[0, 1].set_xlabel('Voltage (V)', fontsize=12, fontweight='bold')
    axes[0, 1].set_ylabel('Current (A)', fontsize=12, fontweight='bold')
    axes[0, 1].set_title(f'I-V Characteristic Curve ({algo_name})', fontsize=14, fontweight='bold')
    axes[0, 1].legend(fontsize=10)
    axes[0, 1].grid(True, alpha=0.3)

    axes[1, 0].plot(V_smooth, P_smooth, 'g-', linewidth=2, label=f'{algo_name} - MTDM')
    axes[1, 0].plot(V_exp, P_exp_arr, 'mo', markersize=6, label='Experimental')
    axes[1, 0].set_xlabel('Voltage (V)', fontsize=12, fontweight='bold')
    axes[1, 0].set_ylabel('Power (W)', fontsize=12, fontweight='bold')
    axes[1, 0].set_title(f'P-V Characteristic Curve ({algo_name})', fontsize=14, fontweight='bold')
    axes[1, 0].legend(fontsize=10)
    axes[1, 0].grid(True, alpha=0.3)

    axes[1, 1].hist(all_rmse, bins=min(20, max(5, n_runs // 2)), color='skyblue', edgecolor='black', alpha=0.7)
    axes[1, 1].axvline(mean_rmse,   color='red',   linestyle='--', linewidth=2, label=f'Mean: {mean_rmse:.4e}')
    axes[1, 1].axvline(median_rmse, color='green', linestyle='--', linewidth=2, label=f'Median: {median_rmse:.4e}')
    axes[1, 1].set_xlabel('RMSE', fontsize=12, fontweight='bold')
    axes[1, 1].set_ylabel('Frequency', fontsize=12, fontweight='bold')
    axes[1, 1].set_title(f'RMSE Distribution ({n_runs} Runs)', fontsize=14, fontweight='bold')
    axes[1, 1].legend(fontsize=10)
    axes[1, 1].grid(True, alpha=0.3, axis='y')

    plt.tight_layout()
    fig1_path = os.path.join(figures_dir, f'{algo_name}_Main.png')
    plt.savefig(fig1_path, dpi=150, bbox_inches='tight')
    plt.close()

    # =========================================================
    # Figure 2 — Current Error
    # =========================================================
    fig2, ax2 = plt.subplots(figsize=(12, 6))
    ax2.plot(range(len(error_current)), error_current, 'o-',
             color='darkred', linewidth=2.5, markersize=8,
             markerfacecolor='coral', markeredgecolor='darkred',
             markeredgewidth=2, label='Current Error')
    ax2.fill_between(range(len(error_current)), error_current, 0,
                     where=(error_current > 0),  alpha=0.3, color='lightcoral', interpolate=True, label='Positive')
    ax2.fill_between(range(len(error_current)), error_current, 0,
                     where=(error_current <= 0), alpha=0.3, color='lightblue',  interpolate=True, label='Negative')
    ax2.axhline(0, color='k', linewidth=2)
    ax2.set_xlabel('Data Point Index', fontsize=14, fontweight='bold')
    ax2.set_ylabel('Current Error (A)', fontsize=14, fontweight='bold')
    ax2.set_title(f'{algo_name} — Current Error Distribution (I_exp − I_pred)', fontsize=16, fontweight='bold')
    ax2.legend(fontsize=11)
    ax2.grid(True, alpha=0.3)
    ax2.text(0.02, 0.98,
             f'Max: {max_err:.6e} A\nMAE: {mae:.6e} A\nRMSE: {all_rmse[best_idx]:.6e}',
             transform=ax2.transAxes, fontsize=11, verticalalignment='top',
             bbox=dict(boxstyle='round', facecolor='wheat', alpha=0.8))
    plt.tight_layout()
    fig2_path = os.path.join(figures_dir, f'{algo_name}_CurrentError.png')
    plt.savefig(fig2_path, dpi=150, bbox_inches='tight')
    plt.close()

    # =========================================================
    # Figure 3 — Power Error
    # =========================================================
    mae_p = np.mean(np.abs(error_power))
    fig3, ax3 = plt.subplots(figsize=(12, 6))
    ax3.plot(range(len(error_power)), error_power, 'o-',
             color='darkgreen', linewidth=2.5, markersize=8,
             markerfacecolor='lightgreen', markeredgecolor='darkgreen',
             markeredgewidth=2, label='Power Error')
    ax3.fill_between(range(len(error_power)), error_power, 0,
                     where=(error_power > 0),  alpha=0.3, color='lightgreen', interpolate=True, label='Positive')
    ax3.fill_between(range(len(error_power)), error_power, 0,
                     where=(error_power <= 0), alpha=0.3, color='lightyellow', interpolate=True, label='Negative')
    ax3.axhline(0, color='k', linewidth=2)
    ax3.set_xlabel('Data Point Index', fontsize=14, fontweight='bold')
    ax3.set_ylabel('Power Error (W)', fontsize=14, fontweight='bold')
    ax3.set_title(f'{algo_name} — Power Error Distribution (P_exp − P_pred)', fontsize=16, fontweight='bold')
    ax3.legend(fontsize=11)
    ax3.grid(True, alpha=0.3)
    ax3.text(0.02, 0.98,
             f'Max: {np.max(np.abs(error_power)):.6e} W\nMAE: {mae_p:.6e} W\nRMSE: {np.sqrt(np.mean(error_power**2)):.6e}',
             transform=ax3.transAxes, fontsize=11, verticalalignment='top',
             bbox=dict(boxstyle='round', facecolor='lightcyan', alpha=0.8))
    plt.tight_layout()
    fig3_path = os.path.join(figures_dir, f'{algo_name}_PowerError.png')
    plt.savefig(fig3_path, dpi=150, bbox_inches='tight')
    plt.close()

    # =========================================================
    # Figure 4 — RMSE Boxplot
    # =========================================================
    fig4, ax4 = plt.subplots(figsize=(10, 8))
    bp = ax4.boxplot([all_rmse], patch_artist=True, notch=False, vert=True, widths=0.5)
    for patch in bp['boxes']:
        patch.set_facecolor('lightcoral')
        patch.set_edgecolor('darkred')
        patch.set_linewidth(2.5)
    for w in bp['whiskers']:
        w.set(color='darkred', linewidth=2)
    for c in bp['caps']:
        c.set(color='darkred', linewidth=2.5)
    for m in bp['medians']:
        m.set(color='blue', linewidth=3)
    for f in bp['fliers']:
        f.set(marker='o', color='red', alpha=0.6, markersize=8)
    ax4.axhline(mean_rmse,   color='green', linestyle='--', linewidth=2, label=f'Mean: {mean_rmse:.4e}')
    ax4.axhline(median_rmse, color='blue',  linestyle=':',  linewidth=2, label=f'Median: {median_rmse:.4e}')
    ax4.set_ylabel('RMSE', fontsize=14, fontweight='bold')
    ax4.set_title(
        f'{algo_name} — RMSE Distribution ({n_runs} Independent Runs)\n'
        f'Min: {min_rmse:.4e}  |  Max: {max_rmse:.4e}  |  Std: {std_rmse:.4e}',
        fontsize=14, fontweight='bold'
    )
    ax4.legend(fontsize=12)
    ax4.grid(True, alpha=0.3, axis='y')
    ax4.set_xticks([1])
    ax4.set_xticklabels([f'{algo_name} RMSE'], fontsize=12, fontweight='bold')
    plt.tight_layout()
    fig4_path = os.path.join(figures_dir, f'{algo_name}_Boxplot.png')
    plt.savefig(fig4_path, dpi=150, bbox_inches='tight')
    plt.close()

    print(f"\n  4 Figures saved to '{figures_dir}/':")
    print(f"    - {algo_name}_Main.png")
    print(f"    - {algo_name}_CurrentError.png")
    print(f"    - {algo_name}_PowerError.png")
    print(f"    - {algo_name}_Boxplot.png")

    # =========================================================
    # Excel Export
    # =========================================================
    excel_filename = os.path.join(excel_dir, f'{algo_name}_Results.xlsx')
    save_comprehensive_excel(
        all_convergence = all_convergence,
        all_rmse        = all_rmse,
        all_best_params = all_best_params,
        V_exp           = V_exp,
        I_exp           = I_exp,
        best_parameters = best_parameters,
        max_iter        = max_iter,
        n_runs          = n_runs,
        all_times       = all_times,
        algo_name       = algo_name,
        full_name       = full_name,
        filename        = excel_filename,
    )

    print("\n" + "=" * 70)
    print(f"  {algo_name} BENCHMARK COMPLETE!")
    print("=" * 70 + "\n")

    return {
        'algo_name': algo_name,
        'full_name': display_title,
        'min_score': min_rmse,
        'min_rmse': min_rmse,
        'mean_rmse': mean_rmse,
        'max_rmse': max_rmse,
        'median_rmse': median_rmse,
        'std_rmse': std_rmse,
        'variance_rmse': variance_rmse,
        'best_parameters': best_parameters,
        'best_convergence': best_convergence,
        'all_rmse': all_rmse,
        'all_times': all_times,
        'all_convergence': all_convergence,
        'excel_path': excel_filename,
        'figures_dir': figures_dir,
    }


# Convenience aliases
run_mtdm_experiment = run_tdm_experiment
run_experiment = run_tdm_experiment
